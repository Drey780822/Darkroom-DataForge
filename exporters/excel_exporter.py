from __future__ import annotations
from pathlib import Path
from typing import List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from models.dataset import Dataset
from .base import BaseExporter


class ExcelExporter(BaseExporter):
    """Exports datasets to formatted Microsoft Excel (.xlsx) workbooks with Wits–merSETA styling."""

    def __init__(self):
        super().__init__(name="excel", extension="xlsx")

    def export_dataset(
        self,
        dataset: Dataset,
        output_path: str,
        include_provenance: bool = False,
        prefer_normalized: bool = True
    ) -> str:
        return self.export_all([dataset], str(Path(output_path).parent), include_provenance=include_provenance)[0]

    def export_all(
        self,
        datasets: List[Dataset],
        output_directory: str,
        include_provenance: bool = False,
        prefer_normalized: bool = True
    ) -> List[str]:
        out_dir = Path(output_directory).resolve()
        out_dir.mkdir(parents=True, exist_ok=True)
        workbook_path = out_dir / "extracted_datasets.xlsx"

        wb = openpyxl.Workbook()
        if not datasets:
            ws = wb.active
            ws.title = "Empty"
            wb.save(str(workbook_path))
            return [str(workbook_path)]

        # Remove default sheet
        wb.remove(wb.active)

        navy_fill = PatternFill(start_color="1D3557", end_color="1D3557", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Segoe UI", size=10)
        thin_border = Border(
            left=Side(style='thin', color="E2E8F0"),
            right=Side(style='thin', color="E2E8F0"),
            top=Side(style='thin', color="E2E8F0"),
            bottom=Side(style='thin', color="E2E8F0")
        )

        for ds in datasets:
            # Sheet names max 31 chars
            sheet_title = ds.metadata.name[:31]
            ws = wb.create_sheet(title=sheet_title)

            df = ds.to_pandas(prefer_normalized=prefer_normalized, include_provenance=include_provenance)

            # Write header
            for col_idx, col_name in enumerate(df.columns, 1):
                cell = ws.cell(row=1, column=col_idx, value=col_name)
                cell.fill = navy_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border

            # Write data rows
            for row_idx, row_values in enumerate(df.values, 2):
                for col_idx, val in enumerate(row_values, 1):
                    val_to_write = "" if val is None else str(val)
                    cell = ws.cell(row=row_idx, column=col_idx, value=val_to_write)
                    cell.font = regular_font
                    cell.border = thin_border

            # Auto-adjust column widths
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    try:
                        if cell.value:
                            max_len = max(max_len, len(str(cell.value)))
                    except:
                        pass
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        wb.save(str(workbook_path))
        return [str(workbook_path)]
